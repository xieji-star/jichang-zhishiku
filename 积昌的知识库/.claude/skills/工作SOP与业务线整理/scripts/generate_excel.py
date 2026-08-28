#!/usr/bin/env python3
"""
Excel 表格生成脚本 - 用于工作SOP与业务线整理skill
生成格式规范的 .xlsx 文件

用法：
  python3 generate_excel.py --type sop --output "路径/文件名.xlsx" --data '[...]'
  python3 generate_excel.py --type business --output "路径/文件名.xlsx" --data '[...]'

--data 参数为 JSON 格式的数组字符串，每条记录为一个字典。
"""

import json
import sys
import argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def create_sop_excel(data, output_path):
    """生成 SOP 流程 Excel 表格"""
    wb = Workbook()
    ws = wb.active
    ws.title = "SOP流程"

    # 表头定义
    headers = ["步骤序号", "步骤名称", "具体操作方法", "所需工具/平台", "注意事项", "预计耗时", "负责人", "向谁汇报"]

    # 样式定义
    header_font = Font(name="微软雅黑", bold=True, size=11, color="000000")
    header_fill = PatternFill(start_color="C8DCF0", end_color="C8DCF0", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="微软雅黑", size=10)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 写入表头
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 写入数据
    for row_idx, item in enumerate(data, 2):
        values = [
            item.get("步骤序号", ""),
            item.get("步骤名称", ""),
            item.get("具体操作方法", ""),
            item.get("所需工具/平台", ""),
            item.get("注意事项", ""),
            item.get("预计耗时", ""),
            item.get("负责人", ""),
            item.get("向谁汇报", ""),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_alignment
            cell.border = thin_border

    # 设置列宽
    col_widths = [10, 18, 40, 20, 25, 10, 12, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    wb.save(output_path)
    print(f"✅ SOP Excel 已保存: {output_path}")


def create_business_excel(data, output_path):
    """生成业务线 Excel 表格"""
    wb = Workbook()
    ws = wb.active
    ws.title = "业务线流程"

    # 表头定义
    headers = ["环节序号", "环节名称", "职责描述", "涉及人员", "输入资源", "输出成果", "所需工具", "注意事项", "向谁汇报"]

    # 样式定义
    header_font = Font(name="微软雅黑", bold=True, size=11, color="000000")
    header_fill = PatternFill(start_color="C8DCF0", end_color="C8DCF0", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="微软雅黑", size=10)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 写入表头
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 写入数据
    for row_idx, item in enumerate(data, 2):
        values = [
            item.get("环节序号", ""),
            item.get("环节名称", ""),
            item.get("职责描述", ""),
            item.get("涉及人员", ""),
            item.get("输入资源", ""),
            item.get("输出成果", ""),
            item.get("所需工具", ""),
            item.get("注意事项", ""),
            item.get("向谁汇报", ""),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_alignment
            cell.border = thin_border

    # 设置列宽
    col_widths = [10, 18, 30, 15, 20, 20, 18, 25, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    wb.save(output_path)
    print(f"✅ 业务线 Excel 已保存: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="生成 SOP 或业务线 Excel 表格")
    parser.add_argument("--type", required=True, choices=["sop", "business"], help="表格类型: sop 或 business")
    parser.add_argument("--output", required=True, help="输出文件路径")
    parser.add_argument("--data", required=True, help="JSON 格式的数据数组字符串")
    args = parser.parse_args()

    try:
        data = json.loads(args.data)
    except json.JSONDecodeError as e:
        print(f"❌ JSON 解析失败: {e}", file=sys.stderr)
        sys.exit(1)

    if not isinstance(data, list):
        print("❌ data 必须为 JSON 数组格式", file=sys.stderr)
        sys.exit(1)

    if args.type == "sop":
        create_sop_excel(data, args.output)
    elif args.type == "business":
        create_business_excel(data, args.output)


if __name__ == "__main__":
    main()
