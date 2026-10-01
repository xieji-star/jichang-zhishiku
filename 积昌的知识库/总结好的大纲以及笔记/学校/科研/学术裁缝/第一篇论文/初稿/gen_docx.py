# -*- coding: utf-8 -*-
"""SCI-manuscript md -> docx converter (shared by the method and experiment chapters).

Usage:
    python gen_docx.py --md <chapter.md> --out <chapter.docx> --section 2
"""
import re, argparse, os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn

FONT = 'Times New Roman'
TITLE = ("A Relaxation–Plasticity Coupled Constitutive and Cyclic Interfacial Damage "
         "Framework for the Foldable Reliability Assessment of OLED Panels")
AUTHORS = "First Author¹, Second Author¹, Hui Xu¹*  (placeholders)"
SECTIONS = {1: "Introduction", 2: "Related Work", 3: "Theoretical Model",
            4: "Results and Discussion"}

# ---------------- tokenizer ----------------
BOLD = re.compile(r'\*\*(.+?)\*\*')
HL   = re.compile(r'==(.+?)==')


def _emit(p, text, bold, hl, size):
    if not text:
        return
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = FONT
    if hl:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _parse(p, text, bold, hl, size):
    m_b, m_h = BOLD.search(text), HL.search(text)
    cands = [(m.start(), 'b', m) for m in (m_b,) if m] + \
            [(m.start(), 'h', m) for m in (m_h,) if m]
    if not cands:
        _emit(p, text, bold, hl, size)
        return
    pos, kind, m = min(cands, key=lambda x: x[0])
    if pos > 0:
        _emit(p, text[:pos], bold, hl, size)
    _parse(p, m.group(1), bold or kind == 'b', hl or kind == 'h', size)
    _parse(p, text[m.end():], bold, hl, size)


def add_rich(doc, text, size=10.5, bold=False, align=None, sb=None, sa=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if sb is not None:
        p.paragraph_format.space_before = Pt(sb)
    if sa is not None:
        p.paragraph_format.space_after = Pt(sa)
    _parse(p, text, bold, False, size)
    return p


# ---------------- latex ----------------
FUNCS = r'(sinh|cosh|tanh|sin|cos|tan|exp|log|ln|max|min|det|lim|sup|inf)'
SYM = [(r'\qquad', '  '), (r'\quad', ' '),
       (r'\varepsilon', 'ε'), (r'\epsilon', 'ε'), (r'\sigma', 'σ'),
       (r'\kappa', 'κ'), (r'\delta', 'δ'), (r'\lambda', 'λ'),
       (r'\nu', 'ν'), (r'\tau', 'τ'), (r'\mu', 'μ'), (r'\eta', 'η'),
       (r'\alpha', 'α'), (r'\beta', 'β'), (r'\gamma', 'γ'), (r'\theta', 'θ'),
       (r'\phi', 'φ'), (r'\psi', 'ψ'), (r'\omega', 'ω'), (r'\rho', 'ρ'),
       (r'\pi', 'π'), (r'\Omega', 'Ω'), (r'\Delta', 'Δ'), (r'\Sigma', 'Σ'),
       (r'\sum', 'Σ'), (r'\leq', '≤'), (r'\le', '≤'), (r'\geq', '≥'), (r'\ge', '≥'),
       (r'\times', '×'), (r'\cdot', '·'), (r'\infty', '∞'), (r'\approx', '≈'),
       (r'\to', '→'), (r'\,', ' ')]


def conv_latex(s):
    s = re.sub(r'\\text\{([^}]*)\}', r'\1', s)
    s = re.sub(r'\\bar\{([^}]*)\}', r'\1', s)
    for k, v in SYM:
        s = s.replace(k, v)
    s = re.sub(r'\\' + FUNCS, r'\1', s)
    s = re.sub(r'_\{([^}]*)\}', r'_\1', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s


# ---------------- table ----------------
def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    for ri, cells in enumerate(rows):
        cells = (cells + [''] * ncol)[:ncol]
        row = t.add_row()
        for ci, txt in enumerate(cells):
            p = row.cells[ci].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            _parse(p, txt, ri == 0, False, 9.5)
    return t


# ---------------- doc ----------------
def new_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    return doc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--md', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--section', type=int)
    ap.add_argument('--label', help='override the draft label, e.g. "Related Work"')
    a = ap.parse_args()

    if a.label:
        label = a.label
    else:
        label = "Section %d (%s)" % (a.section, SECTIONS[a.section])

    doc = new_doc()
    add_rich(doc, TITLE, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=10)
    add_rich(doc, AUTHORS, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=2)
    add_rich(doc, "¹ School of Mechanical Engineering, [University]  ·  *Corresponding author  ·  "
                  "Draft — %s" % label,
             size=9, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=14)
    doc.paragraphs[-1].runs[0].italic = True

    lines = open(a.md, encoding='utf-8').read().splitlines()
    i, n = 0, len(lines)
    infence = False
    while i < n:
        raw = lines[i]
        s = raw.strip()

        # fenced blocks (plotting data etc.) are never typeset
        if s.startswith('```'):
            infence = not infence
            i += 1
            continue
        if infence:
            i += 1
            continue

        if not s:
            i += 1
            continue

        # table block
        if s.startswith('|'):
            blk = []
            while i < n and lines[i].strip().startswith('|'):
                blk.append(lines[i].strip())
                i += 1
            rows = []
            for ln in blk:
                if set(ln) <= set('|:- '):
                    continue
                rows.append([c.strip() for c in ln.strip('|').split('|')])
            add_table(doc, rows)
            continue

        # equation
        if s.startswith('$$'):
            inner = s[2:-2] if s.endswith('$$') else s[2:]
            m = re.search(r'\\tag\{([^}]*)\}', inner)
            tag = m.group(1) if m else ''
            if m:
                inner = inner[:m.start()] + inner[m.end():]
            add_rich(doc, '\t' + conv_latex(inner) + '\t(' + tag + ')',
                     size=10.5, sb=4, sa=4)
            i += 1
            continue

        # headings
        if s.startswith('#'):
            lvl = len(s) - len(s.lstrip('#'))
            title = s[lvl:].strip()
            if lvl == 1:
                add_rich(doc, title, size=12, bold=True, sb=0, sa=6)
            elif lvl == 2:
                add_rich(doc, title, size=11, bold=True, sb=10, sa=4)
            else:
                add_rich(doc, title, size=10.5, bold=True, sb=8, sa=4)
            i += 1
            continue

        # caption
        if re.match(r'^(Figure|Table) \d+  ', s):
            add_rich(doc, s, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, sb=6, sa=8)
            i += 1
            continue

        # body
        add_rich(doc, s, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        i += 1

    doc.save(a.out)
    print('saved:', a.out)


if __name__ == '__main__':
    main()
