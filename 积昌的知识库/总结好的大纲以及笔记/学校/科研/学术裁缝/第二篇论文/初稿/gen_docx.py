# -*- coding: utf-8 -*-
"""SCI-manuscript md -> docx converter (Paper #2).

Supports ONE OR MORE chapter .md files, rendered in order into a single Word
document (chapters separated by a page break), so that the manuscript grows
continuously as new chapters are written.

Formulae are typeset with matplotlib mathtext and embedded as transparent PNGs,
so that they appear as proper typeset equations in Word (journal-grade layout):
    body            Times New Roman 10.5 pt, justified
    headings        bold, 12 / 11 / 10.5 pt
    page            US Letter 8.5 x 11 in, 1 in margins
    equation        centred, with the equation number flush right

Usage:
    python gen_docx.py --md <ch1.md> [<ch2.md> ...] --out <manuscript.docx> \
        --label "Materials and Methods  ·  Results and Discussion"
"""
import re, argparse, os, hashlib, shutil
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

plt.rcParams['mathtext.fontset'] = 'stix'

FONT = 'Times New Roman'
TITLE = ("A Biochar-Reinforced Bio-based Superabsorbent Composite with Slow-release "
         "Nitrogen for Water Retention in Water-saving Irrigation")
AUTHORS = "First Author¹, Second Author¹, [Materials Supervisor]¹*, Yong Liu²*  (placeholders)"
AFFIL = ("¹ School of Mechanical Engineering, [University]  ·  "
         "² School of Water Conservancy and Civil Engineering, [University]  ·  "
         "*Corresponding author")
EQ_DPI = 300
EQ_FONTSIZE = 11.0
EQ_MAX_W = 5.0          # inches

# ---------------- tokenizer ----------------
TOK = re.compile(r'(\*\*.+?\*\*|==.+?==|\*.+?\*)')


def _emit(p, text, bold, hl, italic, size):
    if not text:
        return
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = FONT
    if hl:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _parse(p, text, bold=False, hl=False, size=10.5, italic=False):
    pos = 0
    for m in TOK.finditer(text):
        if m.start() > pos:
            _emit(p, text[pos:m.start()], bold, hl, italic, size)
        tok = m.group(0)
        if tok.startswith('**'):
            _parse(p, tok[2:-2], True, hl, size, italic)
        elif tok.startswith('=='):
            _parse(p, tok[2:-2], bold, True, size, italic)
        else:
            _parse(p, tok[1:-1], bold, hl, size, True)
        pos = m.end()
    if pos < len(text):
        _emit(p, text[pos:], bold, hl, italic, size)


def add_rich(doc, text, size=10.5, bold=False, align=None, sb=None, sa=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if sb is not None:
        p.paragraph_format.space_before = Pt(sb)
    if sa is not None:
        p.paragraph_format.space_after = Pt(sa)
    _parse(p, text, bold, False, size)
    return p


# ---------------- equation (mathtext -> png) ----------------
def render_eq(latex, out_png):
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, latex, fontsize=EQ_FONTSIZE)
    fig.savefig(out_png, dpi=EQ_DPI, transparent=True,
                bbox_inches='tight', pad_inches=0.03)
    plt.close(fig)


def add_equation(doc, latex, tag, tmpdir):
    png = os.path.join(tmpdir, hashlib.md5((latex + tag).encode('utf-8')).hexdigest() + '.png')
    render_eq(latex, png)
    w, h = Image.open(png).size
    width_in = w / EQ_DPI
    if width_in > EQ_MAX_W:
        width_in = EQ_MAX_W
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    p.add_run().add_picture(png, width=Inches(width_in))
    if tag:
        r = p.add_run('\t(' + tag + ')')
        r.font.size = Pt(10.5)
        r.font.name = FONT
    return p


# ---------------- table ----------------
def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    for ri, cells in enumerate(rows):
        cells = (cells + [''] * ncol)[:ncol]
        row = t.add_row()
        for ci, txt in enumerate(cells):
            p = row.cells[ci].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            _parse(p, txt, ri == 0, False, 9.5)
    return t


# ---------------- doc ----------------
def new_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    return doc


def render_chapter(doc, md_path, tmpdir):
    """Render one chapter .md file into `doc`."""
    lines = open(md_path, encoding='utf-8').read().splitlines()
    i, n = 0, len(lines)
    infence = False
    while i < n:
        raw = lines[i]
        s = raw.strip()

        if s.startswith('```'):
            infence = not infence
            i += 1
            continue
        if infence:
            i += 1
            continue
        if not s:
            i += 1
            continue

        # table block
        if s.startswith('|'):
            blk = []
            while i < n and lines[i].strip().startswith('|'):
                blk.append(lines[i].strip())
                i += 1
            rows = []
            for ln in blk:
                if set(ln) <= set('|:- '):
                    continue
                rows.append([c.strip() for c in ln.strip('|').split('|')])
            add_table(doc, rows)
            continue

        # equation
        if s.startswith('$$'):
            inner = s[2:-2] if s.endswith('$$') else s[2:]
            m = re.search(r'\\tag\{([^}]*)\}', inner)
            tag = m.group(1) if m else ''
            if m:
                inner = inner[:m.start()] + inner[m.end():]
            inner = inner.strip()
            add_equation(doc, inner, tag, tmpdir)
            i += 1
            continue

        # headings
        if s.startswith('#'):
            lvl = len(s) - len(s.lstrip('#'))
            title = s[lvl:].strip()
            if lvl == 1:
                add_rich(doc, title, size=12, bold=True, sb=0, sa=6)
            elif lvl == 2:
                add_rich(doc, title, size=11, bold=True, sb=10, sa=4)
            else:
                add_rich(doc, title, size=10.5, bold=True, sb=8, sa=4)
            i += 1
            continue

        # caption
        if re.match(r'^(Figure|Table) \d+  ', s):
            add_rich(doc, s, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, sb=6, sa=8)
            i += 1
            continue

        # body
        add_rich(doc, s, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        i += 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--md', nargs='+', required=True,
                    help='one or more chapter .md files, rendered in the given order')
    ap.add_argument('--out', required=True)
    ap.add_argument('--label',
                    default='Materials and Methods  ·  Results and Discussion')
    a = ap.parse_args()

    tmpdir = '_eqimg'
    shutil.rmtree(tmpdir, ignore_errors=True)
    os.makedirs(tmpdir, exist_ok=True)

    doc = new_doc()
    add_rich(doc, TITLE, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=10)
    add_rich(doc, AUTHORS, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=2)
    add_rich(doc, AFFIL + "  ·  Draft — " + a.label,
             size=9, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=14)
    doc.paragraphs[-1].runs[0].italic = True

    for k, md_path in enumerate(a.md):
        if k:                       # start each new chapter on a fresh page
            doc.add_page_break()
        render_chapter(doc, md_path, tmpdir)

    doc.save(a.out)
    shutil.rmtree(tmpdir, ignore_errors=True)
    print('saved:', a.out, '| chapters:', len(a.md))


if __name__ == '__main__':
    main()
